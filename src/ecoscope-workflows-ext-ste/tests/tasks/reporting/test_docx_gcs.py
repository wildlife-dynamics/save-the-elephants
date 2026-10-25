"""The docx report helpers read their inputs from, and save their output to, GCS (ecoscope-web)."""

import io

import pytest
from cloudpathlib.local import LocalGSClient, LocalGSPath
from docx import Document
from docxtpl import DocxTemplate, InlineImage
from ecoscope_workflows_ext_custom.storage import read_bytes, write_bytes
from ecoscope_workflows_ext_custom.tasks.io import _storage
from PIL import Image

from ecoscope_workflows_ext_ste.tasks.reporting._mapbook_context import build_docx_context
from ecoscope_workflows_ext_ste.tasks.reporting._merge_documents import merge_docx_documents

ROOT = "gs://bucket/user/run"


@pytest.fixture
def fake_gcs(tmp_path, monkeypatch):
    """Back gs:// paths with a local directory instead of real GCS."""
    client = LocalGSClient(local_storage_dir=tmp_path / "gcs")
    monkeypatch.setattr(_storage, "_gs_path", lambda uri: LocalGSPath(uri, client=client))
    return client


def _docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(buf, format="PNG")
    return buf.getvalue()


def _text(path: str) -> list[str]:
    return [p.text for p in Document(io.BytesIO(read_bytes(path))).paragraphs if p.text]


def test_merge_docx_documents_on_gcs(fake_gcs):
    cover = write_bytes(_docx_bytes("Cover"), ROOT, "cover.docx")
    pages = [
        write_bytes(_docx_bytes("Page B"), ROOT, "page_b.docx"),
        write_bytes(_docx_bytes("Page A"), ROOT, "page_a.docx"),
    ]

    out = merge_docx_documents(cover_page_path=cover, context_page_items=pages, output_dir=ROOT, filename="report.docx")

    assert out.endswith("report.docx")
    assert _text(out) == ["Cover", "Page A", "Page B"]


def test_merge_docx_documents_skips_missing_page_on_gcs(fake_gcs):
    cover = write_bytes(_docx_bytes("Cover"), ROOT, "cover.docx")
    pages = [write_bytes(_docx_bytes("Page A"), ROOT, "page_a.docx"), f"{ROOT}/missing.docx"]

    out = merge_docx_documents(cover_page_path=cover, context_page_items=pages, output_dir=ROOT, filename="r")

    assert _text(out) == ["Cover", "Page A"]


def test_build_docx_context_embeds_gcs_image(fake_gcs):
    image = write_bytes(_png_bytes(), ROOT, "map.png")
    template = DocxTemplate(io.BytesIO(_docx_bytes("{{ map }}")))

    ctx = build_docx_context({"map": image, "missing": f"{ROOT}/missing.png", "name": "Lion"}, template)

    assert isinstance(ctx["map"], InlineImage)
    assert ctx["missing"] == f"{ROOT}/missing.png"  # not an existing image: passed through as text
    assert ctx["name"] == "Lion"


def test_create_context_page_on_gcs(fake_gcs):
    from ecoscope_workflows_ext_ste.tasks.reporting._cover_context import create_context_page

    template = write_bytes(_docx_bytes("{{ org_logo }} Prepared by {{ prepared_by }}"), ROOT, "cover.docx")
    logo = write_bytes(_png_bytes(), ROOT, "logo.png")

    out = create_context_page(
        template_path=template,
        output_dir=ROOT,
        context={"org_logo_path": logo, "prepared_by": "Lion <Guardians> & Co"},
        filename="cover_page.docx",
    )

    doc = Document(io.BytesIO(read_bytes(out)))
    assert len(doc.element.body.xpath(".//pic:pic")) == 1
    assert "Prepared by Lion <Guardians> & Co" in _text(out)[0]
